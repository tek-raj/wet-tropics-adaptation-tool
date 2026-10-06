"""Build the content review Word document (v2) from the app's content.

Usage:  python make_content_doc.py content.json output.docx
content.json is window.APP_CONTENT exported from app/js/content.js.

Each question table carries a grey "ID:" code so the completed document can be
read back into content.js automatically. Keep IDs in sync with content.js.
"""
import json
import sys
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x1F, 0x5F, 0x4A)
GREY = RGBColor(0x80, 0x80, 0x80)
FILL_TODO = "FFF6CC"   # yellow: please complete
FILL_DONE = "FFFFFF"   # white: already has text, edit if needed
FILL_HEAD = "E8EFEA"
FILL_LABEL = "F4F7F5"

TOTAL_W = Cm(17)       # A4 portrait minus 2 cm margins
LABEL_W = Cm(4.6)

HELP_ROWS = [
    ("instructions", "Instructions", "What should people write in this box?"),
    ("examples", "Examples", "One or more example answers. Put each example on a new line."),
    ("guidance", "Guidance", "Tips, things to consider, common mistakes."),
    ("moreInfo", "More information link", "Full web address of the matching page on the resource website. Leave blank if the page is not ready."),
]


# ---------------- low-level helpers ----------------
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = w


def cant_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)
    if header:
        h = OxmlElement("w:tblHeader")
        trPr.append(h)


def min_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def keep_with_next(par):
    par.paragraph_format.keep_with_next = True


def write_text(cell, text, size=10, bold=False, italic=False, color=None):
    """Replace cell content with text; newlines become separate paragraphs."""
    cell.text = ""
    lines = (text or "").split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = color
    return cell


def add_field(par, instr):
    """Insert a Word field (e.g. PAGE) into a paragraph."""
    r = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r._r.append(b); r._r.append(t); r._r.append(e)
    return r


# ---------------- document pieces ----------------
def heading(doc, text, level, new_page=False):
    h = doc.add_heading(text, level=level)
    # page_break_before avoids blank pages that a separate page-break paragraph can cause
    h.paragraph_format.page_break_before = new_page
    for r in h.runs:
        r.font.color.rgb = GREEN
    keep_with_next(h)
    return h


def para(doc, text, size=10.5, italic=False, color=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def question_table(doc, title, qid, context, help_obj, extra_rows=(), help_rows=HELP_ROWS, after_rows=()):
    """One table per question: header row + label/answer rows."""
    rows = list(extra_rows) + [(k, lab, hint) for k, lab, hint in help_rows] + list(after_rows)
    t = doc.add_table(rows=1 + len(rows), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    head = t.rows[0].cells[0].merge(t.rows[0].cells[1])
    shade(head, FILL_HEAD)
    head.text = ""
    p = head.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = GREEN
    if context:
        p2 = head.add_paragraph()
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(context)
        r2.italic = True
        r2.font.size = Pt(9)
    p3 = head.add_paragraph()
    p3.paragraph_format.space_after = Pt(0)
    r3 = p3.add_run("ID: " + qid)
    r3.font.size = Pt(7.5)
    r3.font.color.rgb = GREY
    cant_split(t.rows[0])

    for i, (key, label, hint) in enumerate(rows, start=1):
        row = t.rows[i]
        lc, vc = row.cells
        shade(lc, FILL_LABEL)
        lc.text = ""
        lp = lc.paragraphs[0]
        lp.paragraph_format.space_after = Pt(1)
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.size = Pt(9.5)
        hp = lc.add_paragraph()
        hp.paragraph_format.space_after = Pt(0)
        hr = hp.add_run(hint)
        hr.font.size = Pt(8)
        hr.font.color.rgb = GREY
        value = (help_obj or {}).get(key, "") or ""
        write_text(vc, value, size=10)
        shade(vc, FILL_DONE if value.strip() else FILL_TODO)
        min_height(row, 1.0)
        cant_split(row)
    set_widths(t, [LABEL_W, TOTAL_W - LABEL_W])
    # keep the table header with the first rows
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                keep_with_next(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def section_intro_table(doc, sec, key):
    extra = [("intro", "Introduction text", "Short text shown at the top of this tab.")]
    h = dict(sec.get("help") or {})
    h["intro"] = sec.get("intro", "")
    question_table(doc, "About this tab (tab-level ⓘ button)", key,
                   "Shown next to the tab heading. Use it for an overview of the whole step.", h, extra)


def field_context(f):
    kind = {"text": "Short text box", "textarea": "Text box", "date": "Date",
            "radio": "Choose one option", "select": "Drop-down list"}.get(f.get("type"), "")
    if f.get("options"):
        kind += ": " + " / ".join(f["options"])
    return kind


def simple_table(doc, headers, rows, widths, fill_cols=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        write_text(c, htxt, size=9.5, bold=True)
        shade(c, FILL_HEAD)
    cant_split(t.rows[0], header=True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            write_text(c, val, size=9.5)
            if ci in fill_cols:
                shade(c, FILL_TODO)
        min_height(t.rows[ri], 1.1)
        cant_split(t.rows[ri])
    set_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ---------------- review content ----------------
# Sections where the draft guidance needs local, cultural or project knowledge.
# (id, title shown in the document, why it needs review)
REVIEW = [
    ("species.importance", "Why is it important?",
     "Cultural values must be described by, or with the permission of, Traditional Owners. Please check the wording, and the advice about sensitive knowledge, with Rainforest Aboriginal partners."),
    ("site.seasonal", "Seasonal calendar notes",
     "The draft uses general Wet Tropics seasons. Traditional Owner seasonal calendars and local timing (breeding, fruiting, fire) should replace or refine the examples."),
    ("risk.exposure", "Exposure (is the hazard present and getting worse?)",
     "Please confirm which time periods (e.g. 2030/2050) and emissions scenarios the project wants people to use from the hazard maps, so plans are consistent."),
    ("risk.impact", "Vulnerability: how is the species impacted?",
     "The examples use published cases (spectacled flying-fox 2018 heatwave, cassowary after cyclones, lemuroid ringtail possum heat limit). Please check them, and add examples for species the project is focusing on."),
    ("risk.priority", "Priority rating",
     "The definitions of Very low to Very high should be agreed project-wide so that ratings mean the same thing in every plan."),
    ("actions.potentialHelp", "Potential actions",
     "Please add or replace examples with actions proven locally (e.g. current programs by Terrain NRM, ranger groups, councils) that people can learn from."),
    ("scoring.cultural", "Culturally acceptable (Y/N)",
     "Wording about cultural authority and free, prior and informed consent should be checked with Traditional Owner partners."),
    ("scoring.cost", "Cost (0–3)",
     "Please set dollar ranges for each score (e.g. 3 = under $5,000) so cost scores are consistent between plans."),
    ("pathway.trigger", "Trigger",
     "Please check the example triggers and suggest local thresholds or monitoring programs people can use (e.g. specific weather stations, survey programs)."),
    ("summary.cultural", "Summary: Cultural considerations",
     "Wording should be agreed with Traditional Owner partners."),
]

# Topics where a page on the resource website would help (question id, suggested page)
RESOURCE_PAGES = [
    ("risk", "How to do a climate risk assessment (exposure × vulnerability), with a worked example"),
    ("risk.exposure", "How to read the hazard maps and interactive climate map for a site"),
    ("risk.impact", "Species attribute table (climate sensitivity of Wet Tropics species)"),
    ("actions", "Catalogue of adaptation actions for Wet Tropics species, with case studies"),
    ("scoring", "Scoring guide, with a worked example of scoring actions"),
    ("scoring.cultural", "Working with Traditional Owners: protocols, consent and cultural knowledge"),
    ("pathway", "Adaptation pathways explained: triggers, turning points and stopping points"),
    ("site.photoPoints", "Photo monitoring method (setting up and re-taking photo points)"),
    ("site.sightings", "How to search for and report species records (WildNet, Atlas of Living Australia)"),
    ("summary", "Example completed adaptation plan and summary page"),
]


def lookup(C, qid):
    """Return (help dict, label) for an ID like 'species.importance' or 'risk'."""
    parts = qid.split(".")
    sec = C[parts[0]]
    if len(parts) == 1:
        return sec.get("help", {}), sec["title"] + " tab (overview)"
    key = parts[1]
    for coll in ("fields", "columns", "criteria", "boxes"):
        for f in sec.get(coll, []):
            if f["id"] == key:
                return f.get("help", {}), f["label"]
    if key == "cultural" and "cultural" in sec:
        return sec["cultural"]["help"], sec["cultural"]["label"]
    if key in sec and isinstance(sec[key], dict):
        return sec[key], key
    raise KeyError(qid)


# ---------------- build ----------------
def build(C, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name in ("Heading 1", "Heading 2", "Title"):
        doc.styles[name].font.name = "Calibri"

    hp = sec.header.paragraphs[0]
    hr = hp.add_run("Adaptation Planning Tool: content review")
    hr.font.size = Pt(8)
    hr.font.color.rgb = GREY
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Page ")
    fr.font.size = Pt(8)
    add_field(fp, "PAGE").font.size = Pt(8)

    # ---- cover
    t = doc.add_paragraph()
    r = t.add_run(C["title"])
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = GREEN
    t2 = doc.add_paragraph()
    r = t2.add_run("Adaptation Planning Tool: content review (version 2)")
    r.font.size = Pt(14)
    para(doc, "For contributors · " + date.today().strftime("%d/%m/%Y"), size=9.5, color=GREY, after=14)

    heading(doc, "What has changed since version 1", 1)
    para(doc, "Thank you for the suggestions on version 1. These have now been added to the tool:")
    bullet(doc, " the interactive climate map is linked as the Visualisation tool (Site tab) and next to the Hazard maps.", "Visualisation tool:")
    bullet(doc, " the CoastAdapt sea level rise information is linked on the Sea level rise hazard.", "Sea level rise:")
    bullet(doc, " until the species attribute table is published on the NESP website, the tool links to the federal and Queensland recovery plan lists.", "Species attribute table:")
    bullet(doc, " a new Plan focus question (single species, group of species, particular hazard, site or property), plus a Purpose of the plan question.", "Start tab:")
    bullet(doc, " the Plan title question now links to Terrain NRM’s climate page.", "Terrain link:")
    para(doc, "Instructions, examples and guidance have also been written for every question in the tool "
              "(more than 45 questions, including a definition for each 0–3 scoring level). Straightforward questions "
              "such as Place name, GPS and Property contact are complete and are not included here.")

    heading(doc, "What we need from you", 1)
    bullet(doc, " the design decisions still open.", "Part A:")
    bullet(doc, " two links still missing, and new links to confirm.", "Part B:")
    bullet(doc, " the resource website pages that would help most, and their addresses once written.", "Part C:")
    bullet(doc, " 10 questions where the draft needs local, cultural or project knowledge. Edit the draft directly, or write in the yellow box.", "Part D:")
    para(doc, "Please don’t change the grey “ID:” codes; we use them to load your text into the tool. "
              "Turn on Track Changes if several people are editing.", size=10, italic=True)

    heading(doc, "Contributors", 2)
    para(doc, "Please add your details so we know who contributed and who to contact about each part: your name, "
              "your organisation, the parts you reviewed (e.g. “Parts A and D”) and the date you finished.", size=10)
    simple_table(doc, ["Name", "Organisation", "Parts reviewed", "Date (DD/MM/YYYY)"], [["", "", "", ""]] * 4,
                 [Cm(4.2), Cm(4.2), Cm(5.2), Cm(3.4)], fill_cols=(0, 1, 2, 3))

    # ---- Part A decisions
    heading(doc, "Part A: Decisions needed on the template", 1, new_page=True)
    para(doc, "The middle column shows what the tool does for now. Please record the group’s decision in the right-hand column.")
    decisions = [
        ["1", "Scoring direction for Cost and Risk of negative consequences. A high score could mean “bad”, which would distort the total out of 21.",
         "3 = most favourable (lowest cost, lowest risk). This is shown clearly in the tool and the reports."],
        ["2", "The summary page needs Actions NOW / LATER, but the template never asks which actions are now and which are later.",
         "Added a “When: Now / Later” choice to each pathway."],
        ["3", "The summary page needs Cultural considerations, but no question collects them.",
         "The summary draft uses “Why is it important?” and any actions marked not culturally acceptable. Should there be a dedicated question, and on which tab?"],
        ["4", "Conservation status: the Qld list has categories the template leaves out (e.g. Near threatened, Special least concern).",
         "One choice per list, each with its own “Not listed”. Guidance tells people to note other Qld categories under “Why is it important?”."],
        ["5", "“Existing monitoring” is asked twice, once on the Species tab and once on the Site tab.",
         "Both kept: “of the species” (anywhere) and “at the site” (any monitoring at this site)."],
        ["6", "The climate hazards overlap: increased temperature, heatwaves, hot days (>35°C) and hot nights (>20°C). Are the list and thresholds right for the Wet Tropics?",
         "All kept as in the template. Guidance tells people to skip hazards that don’t apply."],
    ]
    simple_table(doc, ["#", "Issue", "What the tool does now", "Decision"],
                 [d + [""] for d in decisions], [Cm(0.8), Cm(6.0), Cm(5.2), Cm(5.0)], fill_cols=(3,))

    # ---- Part B links
    heading(doc, "Part B: Links", 1)
    heading(doc, "Still needed", 2)
    L = C["links"]
    simple_table(doc, ["Link", "Where it appears", "Web address"],
                 [["Resource website home page\nID: resourceSiteUrl", "Tool header (“Resource library”)", ""],
                  ["Species attribute table\nID: links.speciesAttributeTable", "Risk assessment → Vulnerability (replaces the recovery plan links)", ""]],
                 [Cm(5.0), Cm(5.5), Cm(6.5)], fill_cols=(2,))
    heading(doc, "Added — please confirm", 2)
    added = [
        ("Interactive climate map / Visualisation tool", "Site tab heading; Risk → Exposure", L.get("climateMap", "")),
        ("CoastAdapt sea level rise information", "Risk → Sea level rise hazard", L.get("seaLevelRise", "")),
        ("Recovery plans (Federal, DCCEEW)", "Risk → Vulnerability; Species → Ecosystem and habitat", L.get("recoveryPlansFederal", "")),
        ("Recovery and conservation plans (Queensland)", "Risk → Vulnerability", L.get("recoveryPlansQld", "")),
        ("Terrain NRM – Climate change: implications for the Wet Tropics", "Start → Plan title", L.get("terrainClimate", "")),
        ("Climate Change in Australia – Wet Tropics", "Risk assessment tab (ⓘ)", L.get("wetTropicsClimate", "")),
        ("Queensland species search (WildNet)", "Site → Sightings of species", L.get("speciesRecords", "")),
        ("CSIRO Indigenous seasonal calendars", "Site → Seasonal calendar notes", L.get("seasonalCalendars", "")),
    ]
    t = simple_table(doc, ["Link", "Where it appears", "Web address", "OK? / better link"],
                     [[a, b, c, ""] for a, b, c in added], [Cm(4.0), Cm(4.0), Cm(5.2), Cm(3.8)], fill_cols=(3,))
    for row in t.rows[1:]:
        for p in row.cells[2].paragraphs:
            for r in p.runs:
                r.font.size = Pt(8)

    # ---- Part C resource pages
    heading(doc, "Part C: Pages for the resource website", 1, new_page=True)
    para(doc, "Each ⓘ panel in the tool has a “More information” link. These are the topics where a page on the "
              "resource website would help people most. Please add the web address when a page exists, and add other topics if needed.")
    rows = []
    for qid, topic in RESOURCE_PAGES:
        _, label = lookup(C, qid)
        rows.append([topic, label + "\nID: " + qid, ""])
    t = simple_table(doc, ["Suggested page", "Linked from (question)", "Web address"], rows,
                     [Cm(7.0), Cm(4.6), Cm(5.4)], fill_cols=(2,))
    for row in t.rows[1:]:
        for r in row.cells[1].paragraphs[-1].runs:
            r.font.size = Pt(7.5)
            r.font.color.rgb = GREY

    # ---- Part D review
    heading(doc, "Part D: Draft guidance needing local knowledge", 1, new_page=True)
    para(doc, "Guidance for these 10 questions has been drafted, but it needs local, cultural or project knowledge to be "
              "right. The white boxes show the draft exactly as it appears in the tool. Edit them directly, or "
              "write changes and extra local examples in the yellow box.")
    review_rows = [("instructions", "Instructions", "What people should write."),
                   ("examples", "Examples", "Sample answers, one per line."),
                   ("guidance", "Guidance", "Tips and things to consider.")]
    notes_row = [("notes", "Your changes or local examples", "Anything to add or change.")]
    for qid, title, why in REVIEW:
        help_obj, _ = lookup(C, qid)
        question_table(doc, title, qid, "Why review: " + why, help_obj,
                       help_rows=review_rows, after_rows=notes_row)

    # ---- Part E comments
    heading(doc, "Part E: Other comments and suggestions", 1, new_page=True)
    para(doc, "Anything else: missing questions, wording changes, other resources to link to.")
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    shade(t.rows[0].cells[0], FILL_TODO)
    min_height(t.rows[0], 9)
    set_widths(t, [TOTAL_W])

    doc.core_properties.title = "Adaptation Planning Tool: content review"
    doc.core_properties.author = "NESP Wet Tropics adaptation planning"
    doc.save(out)


if __name__ == "__main__":
    src, out = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8") as fh:
        build(json.load(fh), out)
    print("wrote", out)
